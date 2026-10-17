"""The package a workspace hands to its accountant.

SyncroFin is not a double-entry ledger, and does not pretend to be: the
accountant keeps the books, and this is the *source material* they book
from. So the deliverable is a set of plain files that open straight in
Excel/LibreOffice on a Brazilian machine, and that any accounting system's
import tool can be pointed at:

    lancamentos.csv          every posted movement in the period
    resumo_por_categoria.csv receitas / despesas per category
    resumo_mensal.csv        receitas / despesas / resultado per month
    plano_de_contas.csv      the category tree (groups + categories)
    contas.csv               accounts and their current balance
    titulos.csv              invoices, receivable and payable, with competência
    LEIA-ME.txt              what each file is and how it was computed

Conventions are the Brazilian spreadsheet ones, chosen so the file needs no
import wizard: `;` as the field separator, `,` as the decimal mark, dates as
dd/mm/aaaa, UTF-8 with a BOM (Excel drops the encoding otherwise and turns
"Alimentação" into "AlimentaÃ§Ã£o"). Amounts are signed: credits positive,
debits negative, so a column sum is a result rather than a puzzle.

Read-only by construction. Nothing here writes to the database.
"""
from __future__ import annotations

import csv
import io
import uuid
import zipfile
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Iterable, Optional, Sequence

import pyzipper
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.account import Account
from app.models.category import Category
from app.models.category_group import CategoryGroup
from app.models.fiscal_emission import FiscalEmission
from app.models.invoice import Invoice
from app.models.transaction import Transaction
from app.models.workspace import Workspace
from app.services import invoice_service

ZERO = Decimal("0")

#: A guard against an export that tries to read the whole database because
#: somebody typed a year 1900. Ten years is more than any statutory
#: retention period asks for.
MAX_PERIOD_DAYS = 3660

_BOM = "﻿"

_DIRECTION_LABEL = {"receivable": "A receber", "payable": "A pagar"}
_STATE_LABEL = {
    "draft": "Rascunho",
    "open": "Em aberto",
    "partial": "Parcialmente paga",
    "paid": "Paga",
    "overdue": "Vencida",
    "void": "Cancelada",
    "uncollectible": "Incobrável",
}
_ACCOUNT_TYPE_LABEL = {
    "checking": "Conta corrente",
    "savings": "Poupança",
    "credit_card": "Cartão de crédito",
    "wallet": "Carteira",
}


class ExportPeriodError(ValueError):
    """The requested period cannot be exported."""


def validate_period(start: date, end: date) -> None:
    if start > end:
        raise ExportPeriodError("A data inicial não pode ser posterior à final.")
    if (end - start).days > MAX_PERIOD_DAYS:
        raise ExportPeriodError("O período máximo de exportação é de 10 anos.")


# ---------------------------------------------------------------------------
# Formatting — the whole reason a file opens cleanly in a pt-BR spreadsheet
# ---------------------------------------------------------------------------
def fmt_money(value: Optional[Decimal]) -> str:
    """`1234.5` -> `1234,50`. No thousands separator: it would end up in the
    same cell as the number and stop Excel reading it as one."""
    if value is None:
        return ""
    return f"{Decimal(value):.2f}".replace(".", ",")


def fmt_date(value: Optional[date]) -> str:
    return value.strftime("%d/%m/%Y") if value else ""


def _csv_bytes(header: Sequence[str], rows: Iterable[Sequence[object]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(header)
    for row in rows:
        writer.writerow(["" if cell is None else cell for cell in row])
    return (_BOM + buf.getvalue()).encode("utf-8")


def _safe_cell(value: Optional[str]) -> str:
    """Neutralise spreadsheet formula injection.

    A description imported from a bank statement is attacker-controlled text,
    and a cell that starts with `=`, `+`, `-` or `@` is executed as a formula
    when the accountant opens the file. Prefixing an apostrophe makes it text.
    """
    if not value:
        return ""
    return "'" + value if value[0] in ("=", "+", "-", "@", "\t", "\r") else value


def _kind(tx: Transaction, category: Optional[Category]) -> str:
    if tx.source == "opening_balance":
        return "Saldo inicial"
    if tx.transfer_pair_id is not None or (category is not None and category.treat_as_transfer):
        return "Transferência"
    return "Receita" if tx.type == "credit" else "Despesa"


def _signed(tx: Transaction) -> Decimal:
    amount = abs(Decimal(tx.amount))
    return amount if tx.type == "credit" else -amount


def _signed_primary(tx: Transaction, primary_currency: str) -> Optional[Decimal]:
    """The amount in the workspace's main currency, when it is knowable."""
    if tx.amount_primary is not None:
        base = abs(Decimal(tx.amount_primary))
    elif tx.currency == primary_currency:
        base = abs(Decimal(tx.amount))
    else:
        return None
    return base if tx.type == "credit" else -base


# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------
async def _load_transactions(
    session: AsyncSession, workspace_id: uuid.UUID, start: date, end: date
) -> list[Transaction]:
    result = await session.execute(
        select(Transaction)
        .where(
            Transaction.workspace_id == workspace_id,
            Transaction.date >= start,
            Transaction.date <= end,
            Transaction.is_ignored.is_(False),
            Transaction.status == "posted",
        )
        .order_by(Transaction.date, Transaction.created_at)
    )
    return list(result.scalars().all())


async def build_files(
    session: AsyncSession,
    workspace: Workspace,
    start: date,
    end: date,
    *,
    today: Optional[date] = None,
) -> dict[str, bytes]:
    """Every file of the package, keyed by name. Pure read."""
    validate_period(start, end)
    today = today or date.today()
    primary = workspace.default_currency or "USD"

    accounts = (
        await session.execute(select(Account).where(Account.workspace_id == workspace.id))
    ).scalars().all()
    account_names = {a.id: (a.display_name or a.name) for a in accounts}

    categories = (
        await session.execute(select(Category).where(Category.workspace_id == workspace.id))
    ).scalars().all()
    groups = (
        await session.execute(
            select(CategoryGroup).where(CategoryGroup.workspace_id == workspace.id)
        )
    ).scalars().all()
    by_category = {c.id: c for c in categories}
    group_names = {g.id: g.name for g in groups}

    transactions = await _load_transactions(session, workspace.id, start, end)

    # -- lancamentos.csv ----------------------------------------------------
    ledger_rows = []
    summary: dict[tuple[str, str], list[Decimal]] = defaultdict(lambda: [ZERO, ZERO])
    monthly: dict[str, list[Decimal]] = defaultdict(lambda: [ZERO, ZERO])
    for tx in transactions:
        category = by_category.get(tx.category_id) if tx.category_id else None
        if category is not None and category.is_ignored:
            continue
        kind = _kind(tx, category)
        signed = _signed(tx)
        signed_primary = _signed_primary(tx, primary)
        group_name = group_names.get(category.group_id, "") if category and category.group_id else ""
        ledger_rows.append(
            [
                fmt_date(tx.date),
                fmt_date(tx.effective_date),
                _safe_cell(account_names.get(tx.account_id, "")),
                _safe_cell(tx.description),
                _safe_cell(tx.payee),
                _safe_cell(category.name if category else "Sem categoria"),
                _safe_cell(group_name),
                kind,
                fmt_money(signed),
                tx.currency,
                fmt_money(signed_primary),
                _safe_cell(tx.notes),
                str(tx.id),
            ]
        )
        # Only income and spending feed the summaries. A transfer moves money
        # between the company's own accounts, and an opening balance is a
        # starting point: counting either would inflate both sides.
        if kind in ("Receita", "Despesa") and signed_primary is not None:
            key = (
                category.name if category else "Sem categoria",
                group_name,
            )
            slot = 0 if signed_primary > 0 else 1
            summary[key][slot] += abs(signed_primary)
            monthly[tx.date.strftime("%Y-%m")][slot] += abs(signed_primary)

    files: dict[str, bytes] = {}
    files["lancamentos.csv"] = _csv_bytes(
        [
            "Data",
            "Data efetiva (caixa)",
            "Conta",
            "Descrição",
            "Favorecido",
            "Categoria",
            "Grupo",
            "Tipo",
            "Valor",
            "Moeda",
            f"Valor em {primary}",
            "Observações",
            "ID",
        ],
        ledger_rows,
    )

    files["resumo_por_categoria.csv"] = _csv_bytes(
        ["Categoria", "Grupo", f"Receitas ({primary})", f"Despesas ({primary})", f"Resultado ({primary})"],
        [
            [
                _safe_cell(name),
                _safe_cell(group),
                fmt_money(income),
                fmt_money(expense),
                fmt_money(income - expense),
            ]
            for (name, group), (income, expense) in sorted(summary.items())
        ],
    )

    files["resumo_mensal.csv"] = _csv_bytes(
        ["Mês", f"Receitas ({primary})", f"Despesas ({primary})", f"Resultado ({primary})"],
        [
            [month, fmt_money(inc), fmt_money(exp), fmt_money(inc - exp)]
            for month, (inc, exp) in sorted(monthly.items())
        ],
    )

    group_by_id = {g.id: g for g in groups}
    files["plano_de_contas.csv"] = _csv_bytes(
        ["Código", "Grupo", "Categoria", "Natureza", "Conta como receita/despesa", "Ignorada nos relatórios"],
        [
            [
                _safe_cell(group_by_id[c.group_id].code if c.group_id and c.group_id in group_by_id else ""),
                _safe_cell(group_names.get(c.group_id, "") if c.group_id else ""),
                _safe_cell(c.name),
                _safe_cell(group_by_id[c.group_id].natureza if c.group_id and c.group_id in group_by_id else ""),
                "Não" if c.treat_as_transfer else "Sim",
                "Sim" if c.is_ignored else "Não",
            ]
            for c in sorted(categories, key=lambda c: (
                group_by_id[c.group_id].code or "" if c.group_id and c.group_id in group_by_id else "",
                group_names.get(c.group_id, ""),
                c.name,
            ))
        ],
    )

    files["contas.csv"] = _csv_bytes(
        ["Conta", "Tipo", "Moeda", "Saldo atual", "Encerrada"],
        [
            [
                _safe_cell(a.display_name or a.name),
                _ACCOUNT_TYPE_LABEL.get(a.type, a.type),
                a.currency,
                fmt_money(a.balance),
                "Sim" if a.is_closed else "Não",
            ]
            for a in sorted(accounts, key=lambda a: (a.display_name or a.name))
        ],
    )

    invoices = (
        await session.execute(
            select(Invoice)
            .where(
                Invoice.workspace_id == workspace.id,
                Invoice.status != "draft",
                Invoice.issue_date >= start,
                Invoice.issue_date <= end,
            )
            .options(selectinload(Invoice.allocations))
            .order_by(Invoice.issue_date, Invoice.number)
        )
    ).scalars().all()
    invoice_ids = [inv.id for inv in invoices]
    emissions_by_invoice: dict[uuid.UUID, FiscalEmission] = {}
    if invoice_ids:
        em_rows = (
            await session.execute(
                select(FiscalEmission)
                .where(
                    FiscalEmission.workspace_id == workspace.id,
                    FiscalEmission.invoice_id.in_(invoice_ids),
                    FiscalEmission.status == "issued",
                )
            )
        ).scalars().all()
        for em in em_rows:
            emissions_by_invoice.setdefault(em.invoice_id, em)
    files["titulos.csv"] = _csv_bytes(
        [
            "Direção",
            "Número",
            "Série",
            "Documento externo",
            "Favorecido",
            "Emissão",
            "Competência",
            "Vencimento",
            "Moeda",
            "Total",
            "Quitado",
            "Saldo",
            "Situação",
            "Dias de atraso",
            "NF-e número",
            "NF-e série",
            "Chave de acesso",
        ],
        [
            [
                _DIRECTION_LABEL.get(inv.direction, inv.direction),
                inv.number,
                _safe_cell(inv.series),
                _safe_cell(inv.external_number),
                _safe_cell(inv.payee.name if inv.payee else ""),
                fmt_date(inv.issue_date),
                fmt_date(inv.competence_date or inv.issue_date),
                fmt_date(inv.due_date),
                inv.currency,
                fmt_money(inv.total),
                fmt_money(invoice_service.allocated_total(inv)),
                fmt_money(invoice_service.balance(inv)),
                _STATE_LABEL.get(invoice_service.derive_state(inv, today), ""),
                invoice_service.days_overdue(inv, today),
                emissions_by_invoice[inv.id].nfe_number if inv.id in emissions_by_invoice else "",
                emissions_by_invoice[inv.id].nfe_serie if inv.id in emissions_by_invoice else "",
                emissions_by_invoice[inv.id].chave_acesso if inv.id in emissions_by_invoice else "",
            ]
            for inv in invoices
        ],
    )

    files["LEIA-ME.txt"] = _readme(workspace, start, end, primary, len(ledger_rows)).encode("utf-8")
    return files


def build_xlsx(files: dict[str, bytes]) -> bytes:
    """Bundle the CSV files into a single XLSX workbook, one sheet per file.

    The sheet names mirror the CSV filenames (without .csv). This alternative
    to the ZIP is handier for accountants who want everything in one Excel
    workbook, importable in Omie, Conta Azul or any XLSX-aware tool.
    """
    try:
        import openpyxl  # type: ignore[import-untyped]
        from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for XLSX export") from exc

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove the blank default sheet

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    for filename, payload in files.items():
        if not filename.endswith(".csv"):
            continue
        sheet_name = filename[:-4][:31]  # Excel max sheet name = 31 chars
        ws = wb.create_sheet(title=sheet_name)
        content = payload.decode("utf-8").lstrip("﻿")  # strip BOM
        reader = csv.reader(io.StringIO(content), delimiter=";")
        for r_idx, row in enumerate(reader, start=1):
            for c_idx, cell_val in enumerate(row, start=1):
                # Try to convert money strings (e.g. "1234,50") to float for Excel
                cell = ws.cell(row=r_idx, column=c_idx)
                if r_idx == 1:
                    cell.value = cell_val
                    cell.font = header_font
                    cell.fill = header_fill
                else:
                    try:
                        cell.value = float(cell_val.replace(".", "").replace(",", "."))
                    except (ValueError, AttributeError):
                        cell.value = cell_val

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _readme(workspace: Workspace, start: date, end: date, primary: str, rows: int) -> str:
    return (
        f"SyncroFin - pacote para o contador\r\n"
        f"Empresa/workspace: {workspace.name}\r\n"
        f"Período: {fmt_date(start)} a {fmt_date(end)}\r\n"
        f"Moeda principal: {primary}\r\n"
        f"Lançamentos no período: {rows}\r\n"
        f"\r\n"
        f"Formato: CSV separado por ponto e vírgula, decimal com vírgula, datas dd/mm/aaaa,\r\n"
        f"UTF-8 com BOM (abre direto no Excel/LibreOffice).\r\n"
        f"\r\n"
        f"lancamentos.csv          Movimentos efetivados (status 'posted') no período, por data\r\n"
        f"                         da transação. Valor com sinal: receita positiva, despesa\r\n"
        f"                         negativa. 'Data efetiva' é a data de caixa (para cartão de\r\n"
        f"                         crédito, o vencimento da fatura). Lançamentos ignorados e\r\n"
        f"                         categorias marcadas como 'ignorar' ficam de fora.\r\n"
        f"resumo_por_categoria.csv Receitas e despesas por categoria. Transferências entre contas\r\n"
        f"                         e saldos iniciais NÃO entram (não são receita nem despesa).\r\n"
        f"resumo_mensal.csv        Mesmo critério, por mês.\r\n"
        f"plano_de_contas.csv      Categorias e grupos usados no controle.\r\n"
        f"contas.csv               Contas e saldo ATUAL (não o saldo ao fim do período).\r\n"
        f"titulos.csv              Faturas emitidas no período (a receber e a pagar), com data\r\n"
        f"                         de competência, vencimento, valor quitado e situação.\r\n"
        f"\r\n"
        f"Observações: os valores em moeda estrangeira só aparecem no resumo quando há a\r\n"
        f"conversão para a moeda principal registrada na própria transação. O SyncroFin é um\r\n"
        f"controle financeiro e não substitui a escrituração contábil.\r\n"
    )


# ---------------------------------------------------------------------------
# Packaging
# ---------------------------------------------------------------------------
def build_zip(files: dict[str, bytes], password: Optional[str] = None) -> bytes:
    """Zip the package; AES-256 when a password is given (same scheme as the
    workspace backup, so the accountant opens it with any ordinary tool)."""
    buf = io.BytesIO()
    if password:
        archive = pyzipper.AESZipFile(
            buf, "w", compression=zipfile.ZIP_DEFLATED, encryption=pyzipper.WZ_AES
        )
        archive.setpassword(password.encode("utf-8"))
    else:
        archive = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED)
    with archive as zf:
        for name, payload in files.items():
            zf.writestr(name, payload)
    return buf.getvalue()
