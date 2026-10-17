import uuid
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_async_session
from app.core.workspace_context import WorkspaceContext, current_workspace, require_permission
from app.models.workspace import Workspace
from app.schemas.dre import DfcReport, DreAnnualReport, DreQuarterlyReport, DreReport
from app.schemas.report import ReportResponse
from app.services import report_service
from app.services.report_pdf import ReportPDFBuilder

# Export endpoints send financial data outside the workspace: require export.read.
_export_gate = require_permission("export.read")

router = APIRouter(prefix="/api/reports", tags=["reports"])


def _financial_year_start_month(tax_jurisdiction: str | None) -> int:
    """Return the first month of the workspace's financial year."""
    return 4 if (tax_jurisdiction or "").upper() == "IN" else 1


def _reject_unsupported_fiscal_year_report(
    period: str | None, interval: str, financial_year_start_month: int
) -> None:
    if period == "ytd" and interval == "yearly" and financial_year_start_month != 1:
        raise HTTPException(
            status_code=422,
            detail="Yearly YTD reports are not supported for non-calendar financial years",
        )


@router.get("/net-worth", response_model=ReportResponse)
async def get_net_worth(
    months: int = Query(12, ge=1, le=24),
    interval: str = Query("monthly", pattern="^(daily|weekly|monthly|yearly)$"),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    asset_group_ids: Optional[list[uuid.UUID]] = Query(None),
    period: str | None = Query(None, pattern="^ytd$"),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    financial_year_start_month = _financial_year_start_month(ctx.workspace.tax_jurisdiction)
    _reject_unsupported_fiscal_year_report(period, interval, financial_year_start_month)
    return await report_service.get_net_worth_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        account_ids=account_ids, asset_group_ids=asset_group_ids, period=period,
        financial_year_start_month=financial_year_start_month,
    )


@router.get("/income-expenses", response_model=ReportResponse)
async def get_income_expenses(
    months: int = Query(12, ge=1, le=24),
    interval: str = Query("monthly", pattern="^(daily|weekly|monthly|yearly)$"),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    period: str | None = Query(None, pattern="^ytd$"),
    days: Optional[int] = Query(None, ge=1, le=730),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    """`days` overrides `months` with an exact rolling window ending today."""
    financial_year_start_month = _financial_year_start_month(ctx.workspace.tax_jurisdiction)
    _reject_unsupported_fiscal_year_report(period, interval, financial_year_start_month)
    return await report_service.get_income_expenses_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        account_ids=account_ids, period=period, days=days,
        financial_year_start_month=financial_year_start_month,
    )


@router.get("/cash-flow", response_model=ReportResponse)
async def get_cash_flow(
    months: int = Query(6, ge=1, le=12),
    interval: str = Query("daily", pattern="^(daily|weekly|monthly)$"),
    baseline: bool = Query(False),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    return await report_service.get_cash_flow_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        baseline=baseline, account_ids=account_ids,
    )


@router.get("/dre", response_model=DreReport)
async def get_dre(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    cost_center_id: uuid.UUID | None = Query(None),
    accounting_mode: str = Query("competence", pattern="^(competence|cash)$"),
    show_budget: bool = Query(False),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    return await report_service.get_dre(
        session, ctx.workspace.id, ctx.user_id, year, month,
        cost_center_id=cost_center_id,
        accounting_mode=accounting_mode,
        show_budget=show_budget,
    )


@router.get("/dre/quarterly", response_model=DreQuarterlyReport)
async def get_dre_quarterly(
    year: int = Query(..., ge=2000, le=2100),
    quarter: int = Query(..., ge=1, le=4),
    cost_center_id: uuid.UUID | None = Query(None),
    accounting_mode: str = Query("competence", pattern="^(competence|cash)$"),
    show_budget: bool = Query(False),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    return await report_service.get_dre_quarterly(
        session, ctx.workspace.id, ctx.user_id, year, quarter,
        cost_center_id=cost_center_id,
        accounting_mode=accounting_mode,
        show_budget=show_budget,
    )


@router.get("/dre/annual", response_model=DreAnnualReport)
async def get_dre_annual(
    year: int = Query(..., ge=2000, le=2100),
    cost_center_id: uuid.UUID | None = Query(None),
    accounting_mode: str = Query("competence", pattern="^(competence|cash)$"),
    show_budget: bool = Query(False),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    return await report_service.get_dre_annual(
        session, ctx.workspace.id, ctx.user_id, year,
        cost_center_id=cost_center_id,
        accounting_mode=accounting_mode,
        show_budget=show_budget,
    )


@router.get("/dfc", response_model=DfcReport)
async def get_dfc(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    cost_center_id: uuid.UUID | None = Query(None),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    return await report_service.get_dfc(
        session, ctx.workspace.id, ctx.user_id, year, month,
        cost_center_id=cost_center_id,
    )


@router.get("/dfc/range", response_model=list[DfcReport])
async def get_dfc_range(
    year: int = Query(..., ge=2000, le=2100),
    month_start: int = Query(..., ge=1, le=12),
    month_end: int = Query(..., ge=1, le=12),
    cost_center_id: uuid.UUID | None = Query(None),
    ctx: WorkspaceContext = Depends(current_workspace),
    session: AsyncSession = Depends(get_async_session),
):
    """DFC para um intervalo de meses (ex.: trimestre ou ano completo)."""
    if month_end < month_start:
        raise HTTPException(status_code=422, detail="month_end must be >= month_start")
    reports = []
    for m in range(month_start, month_end + 1):
        reports.append(
            await report_service.get_dfc(
                session, ctx.workspace.id, ctx.user_id, year, m,
                cost_center_id=cost_center_id,
            )
        )
    return reports


@router.get("/dre/export")
async def export_dre(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    fmt: str = Query("xlsx", alias="format", pattern="^xlsx$"),
    accounting_mode: str = Query("competence", pattern="^(competence|cash)$"),
    show_budget: bool = Query(False),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o DRE do mês como XLSX."""
    from io import BytesIO

    import openpyxl
    from fastapi.responses import Response

    report = await report_service.get_dre(
        session, ctx.workspace.id, ctx.user_id, year, month,
        accounting_mode=accounting_mode,
        show_budget=show_budget,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"DRE {month:02d}-{year}"
    ws.append(["Demonstrativo de Resultado — " + f"{month:02d}/{year}"])
    ws.append([])
    ws.append(["Receita Bruta", float(report.gross_revenue)])
    ws.append(["Deduções", float(report.deductions)])
    ws.append(["Receita Líquida", float(report.net_revenue)])
    ws.append(["Despesas Operacionais", float(report.operating_expenses)])
    ws.append(["Resultado", float(report.result)])
    ws.append([])
    for section in report.sections:
        ws.append([section.label])
        for grp in section.groups:
            ws.append(["  " + grp.group_name, float(grp.subtotal)])
            for line in grp.lines:
                ws.append(["    " + line.label, float(line.amount)])
        ws.append([])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"dre_{year}_{month:02d}.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/dfc/export")
async def export_dfc(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    fmt: str = Query("xlsx", alias="format", pattern="^xlsx$"),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o DFC do mês como XLSX."""
    from io import BytesIO

    import openpyxl
    from fastapi.responses import Response

    report = await report_service.get_dfc(
        session, ctx.workspace.id, ctx.user_id, year, month,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"DFC {month:02d}-{year}"
    ws.append(["Demonstrativo de Fluxo de Caixa — " + f"{month:02d}/{year}"])
    ws.append([])
    ws.append(["Saldo de Abertura", float(report.opening_balance)])
    ws.append([])
    ws.append(["Período", "Entradas", "Saídas", "Saldo"])
    for p in report.periods:
        ws.append([p.date, float(p.inflows), float(p.outflows), float(p.balance)])
    ws.append([])
    ws.append(["Saldo de Fechamento", float(report.closing_balance)])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"dfc_{year}_{month:02d}.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Branded PDF + additional XLSX exports (G2)
# ---------------------------------------------------------------------------


async def _workspace_with_logo(session, workspace_id):
    """Load workspace and its logo bytes (if any) for branded reports."""
    workspace = await session.get(Workspace, workspace_id)
    logo_bytes = None
    if workspace and workspace.logo_storage_key:
        try:
            from app.providers import get_storage_provider

            storage = get_storage_provider()
            logo_bytes = await storage.download(workspace.logo_storage_key)
        except Exception:
            pass
    return workspace, logo_bytes


@router.get("/dre/pdf")
async def export_dre_pdf(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    accounting_mode: str = Query("competence", pattern="^(competence|cash)$"),
    show_budget: bool = Query(False),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o DRE do mês como PDF com branding do workspace."""
    report = await report_service.get_dre(
        session, ctx.workspace.id, ctx.user_id, year, month,
        accounting_mode=accounting_mode,
        show_budget=show_budget,
    )
    workspace, logo_bytes = await _workspace_with_logo(session, ctx.workspace.id)
    builder = ReportPDFBuilder(workspace, logo_bytes)
    pdf = builder.build_dre_pdf(report, currency=ctx.user.primary_currency)
    filename = f"dre_{year}_{month:02d}.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/dfc/pdf")
async def export_dfc_pdf(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    cost_center_id: uuid.UUID | None = Query(None),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o DFC do mês como PDF com branding do workspace."""
    report = await report_service.get_dfc(
        session, ctx.workspace.id, ctx.user_id, year, month,
        cost_center_id=cost_center_id,
    )
    workspace, logo_bytes = await _workspace_with_logo(session, ctx.workspace.id)
    builder = ReportPDFBuilder(workspace, logo_bytes)
    pdf = builder.build_dfc_pdf(report, currency=ctx.user.primary_currency)
    filename = f"dfc_{year}_{month:02d}.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/net-worth/export")
async def export_net_worth_xlsx(
    months: int = Query(12, ge=1, le=24),
    interval: str = Query("monthly", pattern="^(daily|weekly|monthly|yearly)$"),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    period: str | None = Query(None, pattern="^ytd$"),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o relatório de patrimônio líquido como XLSX."""
    from io import BytesIO

    import openpyxl

    financial_year_start_month = _financial_year_start_month(ctx.workspace.tax_jurisdiction)
    _reject_unsupported_fiscal_year_report(period, interval, financial_year_start_month)
    report = await report_service.get_net_worth_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        account_ids=account_ids, period=period,
        financial_year_start_month=financial_year_start_month,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patrimônio Líquido"
    ws.append(["Patrimônio Líquido"])
    ws.append([])
    ws.append(["Valor Atual", report.summary.primary_value])
    ws.append(["Variação", report.summary.change_amount])
    ws.append([])
    ws.append(["Composição"])
    for b in report.summary.breakdowns:
        ws.append([b.label, b.value])
    ws.append([])
    ws.append(["Data", "Valor"])
    for pt in report.trend:
        ws.append([pt.date, pt.value])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = "patrimonio_liquido.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/income-expenses/export")
async def export_income_expenses_xlsx(
    months: int = Query(12, ge=1, le=24),
    interval: str = Query("monthly", pattern="^(daily|weekly|monthly|yearly)$"),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    period: str | None = Query(None, pattern="^ytd$"),
    days: Optional[int] = Query(None, ge=1, le=730),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o relatório de receitas/despesas como XLSX."""
    from io import BytesIO

    import openpyxl

    financial_year_start_month = _financial_year_start_month(ctx.workspace.tax_jurisdiction)
    _reject_unsupported_fiscal_year_report(period, interval, financial_year_start_month)
    report = await report_service.get_income_expenses_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        account_ids=account_ids, period=period, days=days,
        financial_year_start_month=financial_year_start_month,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Receitas e Despesas"
    ws.append(["Receitas e Despesas"])
    ws.append([])
    ws.append(["Valor Atual", report.summary.primary_value])
    ws.append(["Variação", report.summary.change_amount])
    ws.append([])
    ws.append(["Composição"])
    for b in report.summary.breakdowns:
        ws.append([b.label, b.value])
    ws.append([])
    ws.append(["Data", "Valor"])
    for pt in report.trend:
        ws.append([pt.date, pt.value])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = "receitas_despesas.xlsx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/net-worth/pdf")
async def export_net_worth_pdf(
    months: int = Query(12, ge=1, le=24),
    interval: str = Query("monthly", pattern="^(daily|weekly|monthly|yearly)$"),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    period: str | None = Query(None, pattern="^ytd$"),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o relatório de patrimônio líquido como PDF com branding."""
    financial_year_start_month = _financial_year_start_month(ctx.workspace.tax_jurisdiction)
    _reject_unsupported_fiscal_year_report(period, interval, financial_year_start_month)
    report = await report_service.get_net_worth_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        account_ids=account_ids, period=period,
        financial_year_start_month=financial_year_start_month,
    )
    workspace, logo_bytes = await _workspace_with_logo(session, ctx.workspace.id)
    # Build data rows from breakdowns
    data_rows = [(b.label, b.value) for b in report.summary.breakdowns]
    builder = ReportPDFBuilder(workspace, logo_bytes)
    pdf = builder.build_summary_pdf(
        "Patrimônio Líquido",
        data_rows,
        currency=ctx.user.primary_currency,
        period=f"{months} meses",
    )
    filename = "patrimonio_liquido.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/income-expenses/pdf")
async def export_income_expenses_pdf(
    months: int = Query(12, ge=1, le=24),
    interval: str = Query("monthly", pattern="^(daily|weekly|monthly|yearly)$"),
    account_ids: Optional[list[uuid.UUID]] = Query(None),
    period: str | None = Query(None, pattern="^ytd$"),
    days: Optional[int] = Query(None, ge=1, le=730),
    ctx: WorkspaceContext = Depends(_export_gate),
    session: AsyncSession = Depends(get_async_session),
):
    """Exporta o relatório de receitas/despesas como PDF com branding."""
    financial_year_start_month = _financial_year_start_month(ctx.workspace.tax_jurisdiction)
    _reject_unsupported_fiscal_year_report(period, interval, financial_year_start_month)
    report = await report_service.get_income_expenses_report(
        session, ctx.workspace.id, ctx.user_id, months, interval, ctx.user.primary_currency,
        account_ids=account_ids, period=period, days=days,
        financial_year_start_month=financial_year_start_month,
    )
    workspace, logo_bytes = await _workspace_with_logo(session, ctx.workspace.id)
    data_rows = [(b.label, b.value) for b in report.summary.breakdowns]
    builder = ReportPDFBuilder(workspace, logo_bytes)
    pdf = builder.build_summary_pdf(
        "Receitas e Despesas",
        data_rows,
        currency=ctx.user.primary_currency,
        period=f"{months} meses",
    )
    filename = "receitas_despesas.pdf"
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
